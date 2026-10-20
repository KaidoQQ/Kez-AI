import zipfile

import pytest

from app.backend.utils import file_tools

SAMPLE_MD = """# Расходы за месяц

## Продукты

| Категория | Сумма |
|---|---|
| Хлеб | 120 |
| Семечки | 85,5 |

## Итоги

- **Всего**: 205,5
- `проверено`

```python
print("кез")
```
"""


@pytest.mark.parametrize("text, expected", [
    ("сохрани в pdf", "pdf"),
    ("сделай таблицу в excel", "xlsx"),
    ("создай презентацию", "pptx"),
    ("экспортируй в word", "docx"),
    ("в docx пожалуйста", "docx"),
    ("запиши в txt", "txt"),
    ("сохрани как markdown", "md"),
    ("сохрани в файл", None),
])
def test_detect_export_format(text, expected):
    assert file_tools.detect_export_format(text) == expected


def test_parse_markdown_blocks():
    blocks = file_tools.parse_markdown_blocks(SAMPLE_MD)
    kinds = [kind for kind, _ in blocks]
    assert kinds[:3] == ["heading", "heading", "table"]
    table = dict(blocks)["table"]
    assert table == [["Категория", "Сумма"], ["Хлеб", "120"], ["Семечки", "85,5"]]
    assert ("bullet", "Всего: 205,5") in blocks
    assert ("code", 'print("кез")') in blocks


def test_plain_text_strips_markdown():
    plain = file_tools.markdown_to_plain_text(SAMPLE_MD)
    assert "#" not in plain and "**" not in plain and "```" not in plain
    assert "Расходы за месяц" in plain


@pytest.mark.parametrize("fmt", ["md", "txt", "pdf", "docx", "xlsx", "pptx"])
def test_export_document_all_formats(qapp, export_dir, fmt):
    path = file_tools.export_document(SAMPLE_MD, fmt)
    assert path.parent == export_dir
    assert path.suffix == f".{fmt}"
    assert path.stat().st_size > 0


def test_xlsx_tables_become_sheets_with_numbers(export_dir):
    from openpyxl import load_workbook
    path = file_tools.export_document(SAMPLE_MD, "xlsx")
    sheet = load_workbook(path)["Продукты"]
    assert sheet["A1"].value == "Категория"
    assert sheet["B2"].value == 120
    assert sheet["B3"].value == 85.5


def test_docx_roundtrip(export_dir):
    path = file_tools.export_document(SAMPLE_MD, "docx")
    text = file_tools.read_docx_file(path)
    assert "Расходы за месяц" in text


def test_pptx_has_slide_per_section(export_dir):
    from pptx import Presentation
    path = file_tools.export_document(SAMPLE_MD, "pptx")
    titles = [slide.shapes.title.text for slide in Presentation(str(path)).slides]
    assert titles == ["Расходы за месяц", "Продукты", "Итоги"]


def test_extract_text_docx_and_unknown(tmp_path):
    docx_path = tmp_path / "t.docx"
    with zipfile.ZipFile(docx_path, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
            "<w:p><w:r><w:t>Привет</w:t></w:r><w:r><w:t> мир</w:t></w:r></w:p></w:body></w:document>",
        )
    assert file_tools.extract_text(docx_path) == "Привет мир"
    assert file_tools.extract_text(tmp_path / "photo.png") is None


def test_read_text_file_cp1251(tmp_path):
    path = tmp_path / "legacy.txt"
    path.write_bytes("Привет, Kez".encode("cp1251"))
    assert file_tools.read_text_file(path) == "Привет, Kez"
