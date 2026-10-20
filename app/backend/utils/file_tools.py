"""Утилиты чтения вложений и экспорта документов (PDF / Markdown / TXT / DOCX / XLSX / PPTX)."""

import re
import time
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from core.logger import logger
from core.constants.file_const import (
    MAX_TEXT_FILE_CHARS,
    TEXT_FILE_EXTENSIONS,
    DOCX_EXTENSIONS,
)

DOCX_NAMESPACE: str = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def read_text_file(path: Path) -> str:
    """Читает текстовый файл, перебирая популярные кодировки. Обрезает до MAX_TEXT_FILE_CHARS."""
    for encoding in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            text: str = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = path.read_bytes().decode("utf-8", errors="replace")

    if len(text) > MAX_TEXT_FILE_CHARS:
        logger.warning(f"Файл {path.name} обрезан до {MAX_TEXT_FILE_CHARS} символов.")
        text = text[:MAX_TEXT_FILE_CHARS] + "\n\n[... файл обрезан ...]"
    return text


def read_docx_file(path: Path) -> str:
    """Извлекает текст абзацев из .docx без сторонних зависимостей (word/document.xml)."""
    with zipfile.ZipFile(path) as archive:
        xml_bytes: bytes = archive.read("word/document.xml")

    root = ElementTree.fromstring(xml_bytes)
    paragraphs: list[str] = []
    for paragraph in root.iter(f"{DOCX_NAMESPACE}p"):
        texts: list[str] = [node.text or "" for node in paragraph.iter(f"{DOCX_NAMESPACE}t")]
        paragraphs.append("".join(texts))

    text: str = "\n".join(paragraphs).strip()
    return text[:MAX_TEXT_FILE_CHARS]


def extract_text(path: Path) -> str | None:
    """
    Возвращает текст файла, если его можно прочитать как текст (txt/код/docx).
    Для бинарных форматов (PDF, изображения, аудио) возвращает None.
    """
    suffix: str = path.suffix.lower()
    if suffix in TEXT_FILE_EXTENSIONS:
        return read_text_file(path)
    if suffix in DOCX_EXTENSIONS:
        return read_docx_file(path)
    return None


def detect_export_format(text: str) -> str | None:
    """Определяет формат экспорта по тексту запроса пользователя."""
    lowered: str = text.lower()
    if re.search(r"\bpdf\b|пдф", lowered):
        return "pdf"
    if re.search(r"\b(xlsx|xls|excel|эксел\w*|иксел\w*)\b|таблиц\w*\s+excel", lowered):
        return "xlsx"
    if re.search(r"\b(pptx|ppt|powerpoint|презентаци\w*|слайд\w*)\b", lowered):
        return "pptx"
    if re.search(r"\b(docx|doc|word|ворд\w*)\b", lowered):
        return "docx"
    if re.search(r"\b(txt|text|текстов\w*)\b", lowered):
        return "txt"
    if re.search(r"\b(md|markdown|маркдаун)\b", lowered):
        return "md"
    return None


def extract_title(markdown_text: str) -> str:
    """Достаёт заголовок документа из первой строки '# ...'."""
    for line in markdown_text.splitlines():
        stripped: str = line.strip()
        if stripped.startswith("#"):
            return stripped.lstrip("#").strip()
        if stripped:
            return stripped[:60]
    return "Kez document"


def get_exports_dir() -> Path:
    """Папка экспорта из пользовательских настроек."""
    from core.settings import settings
    return Path(settings.export_dir).expanduser()


def build_export_path(title: str, extension: str) -> Path:
    """Формирует безопасное уникальное имя файла в папке экспорта."""
    exports_dir: Path = get_exports_dir()
    exports_dir.mkdir(parents=True, exist_ok=True)
    safe_title: str = re.sub(r"[^\w\s-]", "", title, flags=re.UNICODE).strip()
    safe_title = re.sub(r"\s+", "_", safe_title)[:50] or "Kez_document"
    timestamp_str: str = time.strftime("%Y-%m-%d_%H-%M")
    return exports_dir / f"{safe_title}_{timestamp_str}_{uuid.uuid4().hex[:4]}.{extension}"


def strip_inline_markdown(text: str) -> str:
    """Убирает инлайн-разметку (жирный, курсив, код, ссылки) из строки."""
    text = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text


def markdown_to_plain_text(markdown_text: str) -> str:
    """Убирает основную Markdown-разметку для экспорта в .txt."""
    text: str = re.sub(r"```\w*\n?", "", markdown_text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    return strip_inline_markdown(text).strip() + "\n"


def parse_markdown_blocks(markdown_text: str) -> list[tuple[str, object]]:
    """
    Разбирает Markdown на простые блоки:
    ('heading', (level, text)) | ('bullet', text) | ('numbered', text) | ('table', rows) | ('code', text) | ('paragraph', text)
    """
    blocks: list[tuple[str, object]] = []
    lines: list[str] = markdown_text.splitlines()
    i: int = 0
    while i < len(lines):
        line: str = lines[i].rstrip()
        stripped: str = line.strip()

        if stripped.startswith("```"):
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(code_lines)))
        elif stripped.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells: list[str] = [strip_inline_markdown(c.strip()) for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        elif heading := re.match(r"^(#{1,6})\s+(.*)", stripped):
            blocks.append(("heading", (len(heading.group(1)), strip_inline_markdown(heading.group(2)))))
        elif bullet := re.match(r"^[-*+]\s+(.*)", stripped):
            blocks.append(("bullet", strip_inline_markdown(bullet.group(1))))
        elif numbered := re.match(r"^\d+[.)]\s+(.*)", stripped):
            blocks.append(("numbered", strip_inline_markdown(numbered.group(1))))
        elif stripped:
            blocks.append(("paragraph", strip_inline_markdown(stripped)))
        i += 1
    return blocks


def write_pdf(markdown_text: str, path: Path) -> None:
    """Рендерит Markdown в PDF средствами Qt (QTextDocument + QPdfWriter), без внешних библиотек."""
    from PyQt6.QtCore import QMarginsF
    from PyQt6.QtGui import QTextDocument, QPdfWriter, QPageSize, QPageLayout, QFont

    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageMargins(QMarginsF(18, 18, 18, 18), QPageLayout.Unit.Millimeter)
    writer.setTitle(extract_title(markdown_text))
    writer.setCreator("Kez AI")

    document = QTextDocument()
    document.setDefaultFont(QFont("Segoe UI", 11))
    document.setMarkdown(markdown_text)
    document.print(writer)


def write_docx(markdown_text: str, path: Path) -> None:
    """Сохраняет Markdown как Word-документ (заголовки, списки, таблицы, код)."""
    from docx import Document
    from docx.shared import Pt

    document = Document()
    for block_type, value in parse_markdown_blocks(markdown_text):
        if block_type == "heading":
            level, text = value
            document.add_heading(text, level=min(level, 4) if level > 1 else 0)
        elif block_type == "bullet":
            document.add_paragraph(value, style="List Bullet")
        elif block_type == "numbered":
            document.add_paragraph(value, style="List Number")
        elif block_type == "table" and value:
            columns: int = max(len(row) for row in value)
            table = document.add_table(rows=0, cols=columns)
            table.style = "Table Grid"
            for row_index, row in enumerate(value):
                cells = table.add_row().cells
                for col_index, cell_text in enumerate(row):
                    cells[col_index].text = cell_text
                    if row_index == 0:
                        for run in cells[col_index].paragraphs[0].runs:
                            run.bold = True
        elif block_type == "code":
            run = document.add_paragraph().add_run(value)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            document.add_paragraph(value)
    document.save(str(path))


def write_xlsx(markdown_text: str, path: Path) -> None:
    """
    Сохраняет Markdown в Excel: каждая таблица — отдельный лист.
    Если таблиц нет, текст построчно кладётся на один лист.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    blocks = parse_markdown_blocks(markdown_text)
    tables: list[tuple[str, list[list[str]]]] = []
    last_heading: str = extract_title(markdown_text)
    for block_type, value in blocks:
        if block_type == "heading":
            last_heading = value[1]
        elif block_type == "table" and value:
            tables.append((last_heading, value))

    workbook = Workbook()
    workbook.remove(workbook.active)

    if not tables:
        rows = [[value[1] if block_type == "heading" else str(value)] for block_type, value in blocks if block_type != "table"]
        tables = [(extract_title(markdown_text), rows)]

    used_names: set[str] = set()
    for index, (sheet_title, rows) in enumerate(tables, start=1):
        name: str = re.sub(r"[\[\]:*?/\\]", "", sheet_title)[:28] or f"Sheet{index}"
        if name in used_names:
            name = f"{name[:25]}_{index}"
        used_names.add(name)
        sheet = workbook.create_sheet(title=name)

        for row in rows:
            sheet.append([_to_number(cell) for cell in row])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col_index in range(1, sheet.max_column + 1):
            width: int = max(len(str(c.value or "")) for c in sheet[get_column_letter(col_index)])
            sheet.column_dimensions[get_column_letter(col_index)].width = min(max(width + 2, 10), 60)

    workbook.save(str(path))


def _to_number(value: str) -> str | int | float:
    """Превращает числовые строки в числа, чтобы Excel мог по ним считать."""
    candidate: str = value.replace(" ", "").replace(" ", "").replace(",", ".")
    if re.fullmatch(r"-?\d+", candidate):
        return int(candidate)
    if re.fullmatch(r"-?\d+\.\d+", candidate):
        return float(candidate)
    return value


def write_pptx(markdown_text: str, path: Path) -> None:
    """
    Сохраняет Markdown как презентацию: '# ' — титульный слайд,
    каждый '## ' — новый слайд, списки/абзацы — пункты слайда.
    """
    from pptx import Presentation
    from pptx.util import Pt

    presentation = Presentation()
    title_layout = presentation.slide_layouts[0]
    content_layout = presentation.slide_layouts[1]

    title_slide = presentation.slides.add_slide(title_layout)
    title_slide.shapes.title.text = extract_title(markdown_text)
    title_slide.placeholders[1].text = "Kez AI 🐦"

    current_body = None
    for block_type, value in parse_markdown_blocks(markdown_text):
        if block_type == "heading":
            level, text = value
            if level == 1:
                continue
            slide = presentation.slides.add_slide(content_layout)
            slide.shapes.title.text = text
            current_body = slide.placeholders[1].text_frame
            current_body.clear()
            continue

        if current_body is None:
            slide = presentation.slides.add_slide(content_layout)
            slide.shapes.title.text = extract_title(markdown_text)
            current_body = slide.placeholders[1].text_frame
            current_body.clear()

        lines: list[str]
        if block_type == "table":
            lines = [" | ".join(row) for row in value]
        else:
            lines = [str(value)]
        for line in lines:
            paragraph = current_body.paragraphs[0] if not current_body.paragraphs[0].text else current_body.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(18)

    presentation.save(str(path))


def export_document(markdown_text: str, export_format: str) -> Path:
    """Сохраняет Markdown-документ в выбранном формате и возвращает путь к файлу."""
    title: str = extract_title(markdown_text)
    path: Path = build_export_path(title, export_format)

    writers = {
        "pdf": write_pdf,
        "docx": write_docx,
        "xlsx": write_xlsx,
        "pptx": write_pptx,
    }
    if export_format in writers:
        writers[export_format](markdown_text, path)
    elif export_format == "txt":
        path.write_text(markdown_to_plain_text(markdown_text), encoding="utf-8")
    else:
        path.write_text(markdown_text.strip() + "\n", encoding="utf-8")

    logger.info(f"Документ экспортирован: {path}")
    return path
